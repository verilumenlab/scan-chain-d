"""Short simple PPT: what SCDA does (for presenting)."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(r"d:\scan chain dianosis agent")
ASSETS = ROOT / "output" / "ppt_assets"
OUT = ROOT / "output" / "SCDA_What_This_Project_Does.pptx"

TEAL = RGBColor(0x0F, 0x76, 0x6E)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)


def set_run(run, size=18, bold=False, color=SLATE):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def paint_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0xFA, 0xFB, 0xFC)


def add_title(slide, text, top=0.35, size=30):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(12.3), Inches(0.7))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    set_run(r, size=size, bold=True, color=TEAL)


def add_bullets(slide, lines, left=0.55, top=1.3, width=12.2, height=5.0, size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(12)
        r = p.add_run()
        r.text = "•  " + line
        set_run(r, size=size, color=SLATE)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def add_footer(slide, n, total=5):
    box = slide.shapes.add_textbox(Inches(11.3), Inches(7.05), Inches(1.5), Inches(0.3))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = f"{n} / {total}"
    set_run(r, size=11, color=MUTED)


def add_pic(slide, name, left, top, width=None, height=None):
    path = ASSETS / name
    if not path.exists():
        return
    kw = {}
    if width is not None:
        kw["width"] = Inches(width)
    if height is not None:
        kw["height"] = Inches(height)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), **kw)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 Title
    s = prs.slides.add_slide(blank)
    paint_bg(s)
    add_pic(s, "slide_hero_dashboard.png", 0.5, 0.4, width=12.3)
    box = s.shapes.add_textbox(Inches(0.6), Inches(5.5), Inches(12), Inches(1.5))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "Scan Chain Diagnosis Agent"
    set_run(r, size=32, bold=True, color=TEAL)
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "What this project does — in simple words"
    set_run(r, size=18, color=SLATE)
    add_footer(s, 1)
    add_notes(
        s,
        "Say: This project helps chip test engineers understand scan-test failures quickly. "
        "Instead of reading huge raw logs, they get a clear dashboard.",
    )

    # 2 One-sentence + problem
    s = prs.slides.add_slide(blank)
    paint_bg(s)
    add_title(s, "In one sentence")
    box = s.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(12.1), Inches(1.4))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = (
        "This project reads chip scan-test fail logs and shows "
        "where the problem is — on a simple dashboard."
    )
    set_run(r, size=24, bold=True, color=SLATE)

    add_title(s, "The problem it solves", top=3.0, size=24)
    add_bullets(
        s,
        [
            "Factories run scan tests on chips. When a test fails, the log is huge and hard to read.",
            "Engineers need fast answers: which chain failed? which cell? where to check next?",
            "Doing this by hand takes too long. This tool does it automatically.",
        ],
        top=3.7,
        height=2.8,
        size=18,
    )
    add_footer(s, 2)
    add_notes(
        s,
        "Say the bold sentence first. Then explain: scan tests shift data through flip-flop chains. "
        "A fail log lists mismatches. Our tool parses those logs and points to likely bad places.",
    )

    # 3 How it works
    s = prs.slides.add_slide(blank)
    paint_bg(s)
    add_title(s, "How it works (4 simple steps)")
    add_pic(s, "flow_simple.png", 0.5, 1.2, width=12.3)
    add_bullets(
        s,
        [
            "1. Take ATE fail logs (tester output).",
            "2. Diagnosis engine finds bad chains, cells, and breaks.",
            "3. Creates 10 requirement reports + 12 KPI cards.",
            "4. Engineer uses the dashboard to decide what to probe or fix next.",
        ],
        top=4.6,
        height=2.4,
        size=17,
    )
    add_footer(s, 3)
    add_notes(
        s,
        "Walk left to right on the flow. Emphasize: input = logs, output = clear KPIs and charts for action.",
    )

    # 4 What you get
    s = prs.slides.add_slide(blank)
    paint_bg(s)
    add_title(s, "What you get on the dashboard")
    add_pic(s, "slide_kpi_grid.png", 0.4, 1.15, width=6.8)
    add_bullets(
        s,
        [
            "Which scan chains failed",
            "Which cells are likely bad",
            "Possible chain breaks",
            "Shift vs capture fail types",
            "Worst chain to check first",
            "Where to probe on the chip",
            "How confident the answer is",
        ],
        left=7.5,
        top=1.3,
        width=5.3,
        height=5.5,
        size=16,
    )
    add_footer(s, 4)
    add_notes(
        s,
        "Point at KPI cards. Example from our data: 23 failing chains, top chain channel14 with 496 fails, "
        "confidence about 68.7%. Click a card to see charts and tables.",
    )

    # 5 Closing / example
    s = prs.slides.add_slide(blank)
    paint_bg(s)
    add_title(s, "Simple example + takeaway")
    add_bullets(
        s,
        [
            "Example: channel14 fails most often (496 fails) → start there.",
            "Tool points to a likely cell and a probe spot on the die.",
            "Takeaway: less time reading logs, more time fixing real chip issues.",
            "Open the UI at http://localhost:3000 to demo live.",
        ],
        top=1.25,
        height=2.6,
        size=18,
    )
    add_pic(s, "chart_ranking.png", 0.7, 3.9, width=12.0)
    add_footer(s, 5)
    add_notes(
        s,
        "Close with: This is an automated scan-chain diagnosis helper — find failing chains/cells fast, "
        "rank what matters most, and guide the next debug step. Ask for questions.",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
